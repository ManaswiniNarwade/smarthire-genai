import io
from pathlib import Path
from typing import List, Union, Optional

class DocumentLoader:
    """Loads and extracts text from various document formats (PDF, DOCX, TXT, MD)."""

    @staticmethod
    def load_pdf(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
        """Extract text from a PDF file using pypdf or pdfplumber."""
        text_content = []

        if isinstance(file_source, (str, Path)):
            stream = open(file_source, "rb")
            should_close = True
        elif isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
            should_close = False
        else:
            stream = file_source
            should_close = False

        try:
            # Try pypdf
            try:
                import pypdf
                reader = pypdf.PdfReader(stream)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(page_text.strip())
            except Exception:
                # Fallback to pdfplumber
                try:
                    import pdfplumber
                    if hasattr(stream, "seek"):
                        stream.seek(0)
                    with pdfplumber.open(stream) as pdf:
                        for page in pdf.pages:
                            page_text = page.extract_text()
                            if page_text:
                                text_content.append(page_text.strip())
                except Exception as inner_e:
                    raise RuntimeError(f"Failed to extract text from PDF: {inner_e}")
        finally:
            if should_close:
                stream.close()

        extracted = "\n\n".join(text_content).strip()
        if not extracted:
            raise ValueError("No extractable text found in PDF (it might be scanned/image-only).")
        return extracted

    @staticmethod
    def load_docx(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
        """Extract text from a DOCX document."""
        import docx

        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        else:
            stream = file_source

        doc = docx.Document(stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n".join(paragraphs).strip()

    @staticmethod
    def load_text(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
        """Read standard text/markdown files."""
        if isinstance(file_source, (str, Path)):
            with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                return f.read().strip()
        elif isinstance(file_source, bytes):
            return file_source.decode("utf-8", errors="replace").strip()
        elif hasattr(file_source, "read"):
            content = file_source.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="replace").strip()
            return str(content).strip()
        return str(file_source).strip()

    @classmethod
    def load_document(cls, file_source: Union[str, Path, io.BytesIO, bytes], filename: Optional[str] = None) -> str:
        """Auto-detect format by filename or path extension and extract text."""
        name = ""
        if filename:
            name = filename.lower()
        elif isinstance(file_source, (str, Path)):
            name = str(file_source).lower()
        elif hasattr(file_source, "name"):
            name = str(file_source.name).lower()

        if name.endswith(".pdf"):
            return cls.load_pdf(file_source)
        elif name.endswith(".docx") or name.endswith(".doc"):
            return cls.load_docx(file_source)
        else:
            return cls.load_text(file_source)

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 600, chunk_overlap: int = 100) -> List[str]:
        """
        Split text into overlapping chunks using paragraph, sentence, and word boundaries.
        """
        if not text:
            return []

        if len(text) <= chunk_size:
            return [text]

        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = start + chunk_size
            if end >= text_len:
                chunks.append(text[start:].strip())
                break

            # Find nearest natural break (double newline, single newline, period, space)
            break_point = -1
            for separator in ["\n\n", "\n", ". ", " "]:
                idx = text.rfind(separator, start + chunk_overlap, end)
                if idx != -1:
                    break_point = idx + len(separator)
                    break

            if break_point == -1:
                break_point = end

            chunk = text[start:break_point].strip()
            if chunk:
                chunks.append(chunk)

            start = max(start + 1, break_point - chunk_overlap)

        return chunks
